from __future__ import annotations

import argparse
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
import os
import tempfile
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from lxml import etree


DEFAULT_REFERENCE = Path(
    "C:/Users/admin/Desktop/AGTU Structured/2 курс/2 семестр/"
    "Курсовой проект — Разработка профессиональных приложений/Документы/"
    "Документ к курсовому проекту.docx"
)
DEFAULT_OUTPUT = Path(__file__).with_name("M10_шаблон_стиля_АГТУ.docx")

TITLE = (
    "Моделирование и симуляция робота-манипулятора, "
    "выполняющего задачу сортировки объектов по цвету"
)
STUDENT_SHORT = "Шатуев Р.С."
GROUP = "ДИНРб-31"
SUPERVISOR = "Кузнецова В.Ю."
DISCIPLINE = "Моделирование роботов"

CORE_NS = "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
DC_NS = "http://purl.org/dc/elements/1.1/"
DCTERMS_NS = "http://purl.org/dc/terms/"
XSI_NS = "http://www.w3.org/2001/XMLSchema-instance"

TASKS = (
    "Зафиксировать требования и критерии виртуальной проверки.",
    "Спроектировать архитектуру и параметрическую модель манипулятора SCARA.",
    "Реализовать и проверить прямую и обратную кинематику.",
    "Реализовать распознавание цвета и определение координат объектов.",
    "Разработать безопасное планирование траекторий и цикл захвата и сортировки.",
    "Интегрировать систему в симуляцию, провести виртуальные испытания и подготовить документацию.",
)

LITERATURE = (
    "Lynch K. M., Park F. C. Modern Robotics: Mechanics, Planning, and Control. Cambridge University Press, 2017. URL: https://modernrobotics.northwestern.edu/nu-gm-book-resource/ (дата обращения: 27.09.2026).",
    "Google DeepMind. MuJoCo Documentation: Modeling; XML Reference; API Reference. URL: https://mujoco.readthedocs.io/en/stable/ (дата обращения: 27.09.2026).",
    "OpenCV. OpenCV 4.12.0 Documentation: Color Space Conversions; Features2D + Homography. URL: https://docs.opencv.org/4.12.0/d8/d01/group__imgproc__color__conversions.html; https://docs.opencv.org/4.12.0/d7/dff/tutorial_feature_homography.html (дата обращения: 27.09.2026).",
    "NIST. Guide to the SI, Appendix B.9: Factors for units listed by kind of quantity or field of science. URL: https://www.nist.gov/pml/special-publication-811/nist-guide-si-appendix-b-conversion-factors/nist-guide-si-appendix-b9 (дата обращения: 27.09.2026).",
    "Prusa Polymers. Prusament PLA Technical Data Sheet. URL: https://prusament.com/wp-content/uploads/2022/10/PLA_Prusament_TDS_2021_10_EN.pdf (дата обращения: 27.09.2026).",
    "Проектная документация M1–M9: архитектура, SSOT, кинематика, зрение, планирование, FSM, интеграция, протокол и отчет виртуальных испытаний. Исходники и версии сохранены в составе проекта.",
)

SCHEDULE = (
    "M0. Аудит задания и фиксация требований",
    "M1. Архитектура и выбор средств моделирования",
    "M2. Геометрия модели и единая спецификация параметров",
    "M3. Прямая и обратная кинематика",
    "M4. Компьютерное зрение и определение положения объектов",
    "M5. Планирование траекторий и проверка столкновений",
    "M6. Алгоритм сортировки и конечный автомат",
    "M7–M8. Интеграция физической модели, визуализация и видео",
    "M9. Виртуальные испытания и анализ результатов",
    "M10–M12. Записка, презентация и финальная приемка",
)


def preserve_first_run_properties(paragraph: Paragraph):
    if not paragraph.runs or paragraph.runs[0]._r.rPr is None:
        return None
    return deepcopy(paragraph.runs[0]._r.rPr)


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    rpr = preserve_first_run_properties(paragraph)
    paragraph.clear()
    if not text:
        return
    run = paragraph.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)


def unique_cells(table):
    seen = set()
    for row in table.rows:
        for cell in row.cells:
            # Keep the XML element itself in the set. Using id(cell._tc) is
            # unsafe here: python-docx creates short-lived wrappers, and CPython
            # may reuse their ids while this generator is consumed lazily.
            key = cell._tc
            if key not in seen:
                seen.add(key)
                yield cell


def sanitize_package_metadata(docx_path: Path) -> None:
    """Remove inherited reference metadata and refresh document properties."""
    timestamp = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    temporary = None
    try:
        handle, temporary_name = tempfile.mkstemp(
            prefix=f".{docx_path.stem}.", suffix=".tmp", dir=docx_path.parent
        )
        os.close(handle)
        temporary = Path(temporary_name)
        with ZipFile(docx_path, "r") as source, ZipFile(temporary, "w") as target:
            for item in source.infolist():
                payload = source.read(item.filename)
                if item.filename == "docProps/core.xml":
                    root = etree.fromstring(payload)
                    creator = root.find(f"{{{DC_NS}}}creator")
                    if creator is None:
                        creator = etree.SubElement(root, f"{{{DC_NS}}}creator")
                    creator.text = "Шатуев Руслан Саматович"
                    modifier = root.find(f"{{{CORE_NS}}}lastModifiedBy")
                    if modifier is None:
                        modifier = etree.SubElement(root, f"{{{CORE_NS}}}lastModifiedBy")
                    modifier.text = "Шатуев Руслан Саматович"
                    for name in ("created", "modified"):
                        node = root.find(f"{{{DCTERMS_NS}}}{name}")
                        if node is None:
                            node = etree.SubElement(root, f"{{{DCTERMS_NS}}}{name}")
                        node.text = timestamp
                        node.set(f"{{{XSI_NS}}}type", "dcterms:W3CDTF")
                    last_printed = root.find(f"{{{CORE_NS}}}lastPrinted")
                    if last_printed is not None:
                        root.remove(last_printed)
                    revision = root.find(f"{{{CORE_NS}}}revision")
                    if revision is not None:
                        revision.text = "1"
                    payload = etree.tostring(
                        root, xml_declaration=True, encoding="UTF-8", standalone=True
                    )
                elif item.filename == "docProps/app.xml":
                    root = etree.fromstring(payload)
                    stale_fields = {
                        "Company", "Template", "Pages", "Words", "Characters",
                        "Lines", "Paragraphs", "TotalTime", "CharactersWithSpaces",
                    }
                    for child in list(root):
                        local_name = etree.QName(child).localname
                        if local_name in stale_fields:
                            root.remove(child)
                        elif local_name == "Application":
                            child.text = "python-docx"
                    payload = etree.tostring(
                        root, xml_declaration=True, encoding="UTF-8", standalone=True
                    )
                target.writestr(item, payload)
        os.replace(temporary, docx_path)
    finally:
        if temporary is not None and temporary.exists():
            temporary.unlink()


def replace_in_table(table, replacements: dict[str, str]) -> None:
    for cell in unique_cells(table):
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                for old, new in replacements.items():
                    if old in run.text:
                        run.text = run.text.replace(old, new)


def set_cell_value(cell, text: str, *, font_size: float | None = None) -> None:
    paragraph = cell.paragraphs[0]
    set_paragraph_text(paragraph, text)
    for extra in list(cell.paragraphs[1:]):
        extra._element.getparent().remove(extra._element)
    if font_size is not None:
        for run in paragraph.runs:
            run.font.size = Pt(font_size)


def find_body_paragraph(doc, exact: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == exact:
            return paragraph
    raise ValueError(f"Front-matter paragraph not found: {exact}")


def remove_old_body(doc) -> None:
    body = doc.element.body
    children = list(body)
    intro_index = next(
        (
            index
            for index, child in enumerate(children)
            if child.tag == qn("w:p") and Paragraph(child, doc).text.strip() == "ВВЕДЕНИЕ"
        ),
        None,
    )
    if intro_index is None:
        raise ValueError("Reference document has no body start heading 'ВВЕДЕНИЕ'")
    for child in children[intro_index:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def clear_reference_toc(doc) -> None:
    controls = [element for element in doc.element.body if element.tag == qn("w:sdt")]
    toc = next(
        (
            control
            for control in controls
            if "TOC" in "".join(control.itertext()).upper()
        ),
        None,
    )
    if toc is None:
        raise ValueError("Reference document has no structured table-of-contents field")
    content = toc.find(qn("w:sdtContent"))
    if content is None:
        raise ValueError("Reference table-of-contents control has no content slot")
    for child in list(content):
        content.remove(child)


def replace_assignment_fields(doc) -> None:
    paragraphs = doc.paragraphs
    for paragraph in paragraphs:
        text = paragraph.text.strip()
        if text.startswith("Обучающийся\t"):
            if len(paragraph.runs) >= 3:
                paragraph.runs[-1].text = STUDENT_SHORT
        elif text.startswith("Группа\t"):
            if len(paragraph.runs) >= 3:
                paragraph.runs[-1].text = GROUP
        elif text.startswith("Тема\t"):
            if len(paragraph.runs) >= 3:
                paragraph.runs[-1].text = TITLE

    task_heading = find_body_paragraph(doc, "Задачи")
    paragraph_index = next(i for i, paragraph in enumerate(paragraphs) if paragraph._p is task_heading._p)
    task_paragraphs = [
        paragraph
        for paragraph in paragraphs[paragraph_index + 1 :]
        if paragraph.text.strip()
    ]
    task_paragraphs = task_paragraphs[: len(TASKS)]
    if len(task_paragraphs) != len(TASKS):
        raise ValueError("Reference task list does not have the expected six numbered slots")
    for paragraph, task in zip(task_paragraphs, TASKS):
        set_paragraph_text(paragraph, task)

    literature_heading = find_body_paragraph(doc, "Список рекомендуемой литературы")
    paragraph_index = next(i for i, paragraph in enumerate(paragraphs) if paragraph._p is literature_heading._p)
    literature_paragraphs = [
        paragraph
        for paragraph in paragraphs[paragraph_index + 1 :]
        if paragraph.text.strip()
    ][: len(LITERATURE)]
    if len(literature_paragraphs) != len(LITERATURE):
        raise ValueError("Reference literature list does not have six editable slots")
    for paragraph, citation in zip(literature_paragraphs, LITERATURE):
        set_paragraph_text(paragraph, citation)


def replace_calendar(doc) -> None:
    schedule_table = next(
        (
            table
            for table in doc.tables
            if table.rows
            and len(table.rows[0].cells) >= 4
            and "Разделы, темы" in " ".join(" ".join(cell.text.split()) for cell in table.rows[0].cells)
        ),
        None,
    )
    if schedule_table is None or len(schedule_table.rows) != len(SCHEDULE) + 1:
        raise ValueError("Reference calendar table does not have ten data rows")
    for index, (row, task) in enumerate(zip(schedule_table.rows[1:], SCHEDULE), start=1):
        set_cell_value(row.cells[0], str(index), font_size=11)
        set_cell_value(row.cells[1], task, font_size=11)
        set_cell_value(row.cells[2], "")
        set_cell_value(row.cells[3], "")

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("С графиком ознакомлен"):
            set_paragraph_text(paragraph, "С графиком ознакомлен  «____» _________________ 20___ г.")
        elif text.startswith("Сулейманов Г.Б."):
            set_paragraph_text(
                paragraph,
                f"{STUDENT_SHORT} _______________________, обучающийся группы {GROUP}",
            )
        elif text.startswith("Руководитель курсового проекта"):
            set_paragraph_text(
                paragraph,
                f"Руководитель курсового проекта ___________________________ {SUPERVISOR}",
            )


def populate_front_matter(doc) -> None:
    # Title-page fields, keeping the original paragraphs, logo, columns and typography.
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("Астрахань"):
            set_paragraph_text(paragraph, "Астрахань – 2026")
            # Keep the city/year line on the cover page, as in the reference.
            # The current topic needs two title lines, so the reference's 100 pt
            # spacer otherwise pushes this line onto a mostly empty extra page.
            paragraph.paragraph_format.space_before = Pt(35)
            break

    cover = doc.tables[0]
    for cell in unique_cells(cover):
        for paragraph in cell.paragraphs:
            text = paragraph.text.strip()
            if text.startswith("Веб-платформа для обучения"):
                set_paragraph_text(paragraph, "Моделирование и симуляция робота-манипулятора,")
            elif "Ассоциативные операции над множествами" in text:
                set_paragraph_text(paragraph, "выполняющего задачу сортировки объектов по цвету")
            elif text.startswith("по дисциплине"):
                set_paragraph_text(paragraph, f"по дисциплине «{DISCIPLINE}»")
            elif text.startswith("обучающимся группы"):
                set_paragraph_text(paragraph, f"обучающимся группы {GROUP}")
            elif "Сулеймановым Г.Б." in text:
                set_paragraph_text(paragraph, "Шатуевым Р.С.")
            elif "ассистент Тараканов В. Д." in text:
                set_paragraph_text(paragraph, f"Руководитель\n{SUPERVISOR}")
            elif "2025г." in text:
                set_paragraph_text(paragraph, text.replace("2025г.", "20___г."))
            elif "2025" in text:
                set_paragraph_text(paragraph, text.replace("2025", "20___"))

    # The reference reserves an exact-height title row; fit the longer current
    # topic to it while keeping it prominent above the body text.
    for cell in cover.rows[2].cells:
        for paragraph in cell.paragraphs:
            if paragraph.text.strip().startswith(("Моделирование и симуляция", "выполняющего задачу")):
                for run in paragraph.runs:
                    run.font.size = Pt(14)

    replace_assignment_fields(doc)

    # Approval identities from the old example are not current project data.
    assignment_approval = doc.tables[1]
    for cell in unique_cells(assignment_approval):
        for paragraph in cell.paragraphs:
            text = paragraph.text.strip()
            if text in {"д.т.н., профессор", "д.т.н., профессор "}:
                set_paragraph_text(paragraph, "")
            elif "Хоменко" in text:
                set_paragraph_text(paragraph, "______________________________")
            elif "2025" in text:
                set_paragraph_text(paragraph, "«____»__________________20___ г.")

    date_table = next(
        (
            table
            for table in doc.tables
            if len(table.rows) == 1
            and len(table.columns) == 2
            and "Дата получения задания" in " ".join(" ".join(cell.text.split()) for cell in table.rows[0].cells)
        ),
        None,
    )
    if date_table is None:
        raise ValueError("Assignment date table not found")
    set_cell_value(date_table.cell(0, 1), "«_____»___________20___г.\n«_____»___________20___г.")

    signature_table = next(
        (
            table
            for table in doc.tables
            if len(table.rows) == 1
            and len(table.columns) == 3
            and "Руководитель" in " ".join(cell.text for cell in table.rows[0].cells)
            and "Обучающийся" in " ".join(cell.text for cell in table.rows[0].cells)
        ),
        None,
    )
    if signature_table is None:
        raise ValueError("Assignment signature table not found")
    replace_in_table(signature_table, {
        "ассистент": "",
        "Тараканов В. Д.": SUPERVISOR,
        "Сулейманов Г.Б.": STUDENT_SHORT,
        "ДИНРб-21": GROUP,
        "202___": "20___",
    })
    # The date year is split across separate Word runs in the supplied form,
    # so normalize each complete paragraph rather than relying on run-level
    # replacement alone.
    for cell in unique_cells(signature_table):
        for paragraph in cell.paragraphs:
            if "202___" in paragraph.text:
                set_paragraph_text(paragraph, paragraph.text.replace("202___", "20___"))

    # Replace the second approval block on the calendar sheet.
    calendar_approval = next(
        (
            table
            for table in doc.tables
            if len(table.rows) == 2
            and len(table.columns) == 2
            and "К заданию на курсовую работу" in " ".join(" ".join(cell.text.split()) for cell in table.rows[0].cells)
        ),
        None,
    )
    if calendar_approval is None:
        raise ValueError("Calendar approval block not found")
    for cell in unique_cells(calendar_approval):
        for paragraph in cell.paragraphs:
            text = paragraph.text.strip()
            if "д.т.н." in text:
                set_paragraph_text(paragraph, "")
            elif "Хоменко" in text:
                set_paragraph_text(paragraph, "______________________________")
            elif "2025" in text:
                set_paragraph_text(paragraph, "«____»__________________20___г.")
            elif "К заданию на курсовую работу" in text:
                set_paragraph_text(
                    paragraph,
                    f"К заданию на курсовую работу\nпо дисциплине\n«{DISCIPLINE}»",
                )
            elif text == "приложений»":
                set_paragraph_text(paragraph, "")

    replace_calendar(doc)
    clear_reference_toc(doc)

    # Apply this after template substitutions: the source has mixed-script
    # capitals and Word run formatting that may be recreated during edits.
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("ФЕДЕРАЛЬНОЕ") and len(text) > 50:
            for run in paragraph.runs:
                run.font.size = Pt(11)

    old_terms = (
        "Сулейманов", "Тараканов", "Хоменко", "ДИНРб-21", "Веб-платформа",
        "Ассоциативные операции", "Разработка профессиональных приложений", "2025",
    )
    front_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    front_text += "\n" + "\n".join(cell.text for table in doc.tables for cell in unique_cells(table))
    remaining = [term for term in old_terms if term in front_text]
    if remaining:
        raise ValueError(f"Reference-era values remain in front matter: {remaining}")


def prepare(reference: Path, output: Path) -> None:
    if not reference.is_file():
        raise FileNotFoundError(reference)
    doc = Document(reference)
    if len(doc.sections) < 3:
        raise ValueError("Reference does not contain the expected cover/frontmatter/body sections")
    remove_old_body(doc)
    populate_front_matter(doc)
    doc.core_properties.title = TITLE
    doc.core_properties.subject = DISCIPLINE
    doc.core_properties.author = "Шатуев Руслан Саматович"
    doc.core_properties.keywords = "робот-манипулятор; сортировка по цвету; виртуальная симуляция"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    sanitize_package_metadata(output)


def main() -> None:
    parser = argparse.ArgumentParser(description="Create a sanitized M10 front-matter template from the supplied AGTU reference.")
    parser.add_argument("--reference", type=Path, default=DEFAULT_REFERENCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    prepare(args.reference.resolve(), args.output.resolve())
    print(f"Created front-matter template: {args.output.resolve()}")


if __name__ == "__main__":
    main()
